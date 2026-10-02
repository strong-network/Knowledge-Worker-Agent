#!/usr/bin/env python3
# Copyright 2026 Citrix Systems, Inc.
# SPDX-License-Identifier: Apache-2.0
"""Read a .xlsx: sheets, cells, formulas, named ranges.

Returns TEXT. The workbook is never attached to the conversation.

A spreadsheet is usually far too large to dump whole. Start with --stats,
then read only the range you need.

Usage:
  xlsx_extract.py BOOK.xlsx --stats            # sheets, dimensions, headers
  xlsx_extract.py BOOK.xlsx                    # all sheets as TSV (capped)
  xlsx_extract.py BOOK.xlsx --sheet Sales      # one sheet
  xlsx_extract.py BOOK.xlsx --range A1:D20     # a specific range
  xlsx_extract.py BOOK.xlsx --formulas         # show formulas, not values
  xlsx_extract.py BOOK.xlsx --search "ISO 27001"
  xlsx_extract.py BOOK.xlsx --max-rows 500     # raise the row cap
"""
import argparse
import os

DEFAULT_MAX_ROWS = 200


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\n", " ").replace("\t", " ")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--range", dest="rng")
    ap.add_argument("--formulas", action="store_true")
    ap.add_argument("--search")
    ap.add_argument("--stats", action="store_true")
    ap.add_argument("--max-rows", type=int, default=DEFAULT_MAX_ROWS)
    a = ap.parse_args()

    import openpyxl

    # data_only=True gives cached computed values; False gives formulas.
    wb = openpyxl.load_workbook(a.file, data_only=not a.formulas, read_only=False)

    print(f"# {os.path.basename(a.file)}")
    print(f"sheets: {len(wb.sheetnames)} — {', '.join(wb.sheetnames)}")

    try:
        names = list(wb.defined_names.keys())
        if names:
            print(f"named ranges: {', '.join(names[:20])}")
    except Exception:
        pass

    targets = [a.sheet] if a.sheet else wb.sheetnames
    for name in targets:
        if name not in wb.sheetnames:
            print(f"\n!! no sheet named {name!r}")
            continue
        ws = wb[name]
        print(f"\n## {name}  ({ws.max_row} rows x {ws.max_column} cols)")

        if a.stats:
            hdr = [fmt(c.value) for c in ws[1]] if ws.max_row else []
            if any(hdr):
                print("headers: " + " | ".join(h for h in hdr if h))
            continue

        if a.search:
            hits = 0
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None and a.search.lower() in str(c.value).lower():
                        print(f"  {c.coordinate}: {fmt(c.value)}")
                        hits += 1
                        if hits >= 100:
                            print("  ... (stopped at 100 hits)")
                            break
                if hits >= 100:
                    break
            if not hits:
                print("  no matches")
            continue

        rows = ws[a.rng] if a.rng else ws.iter_rows(max_row=min(ws.max_row, a.max_rows))
        n = 0
        for row in rows:
            cells = row if isinstance(row, tuple) else (row,)
            line = "\t".join(fmt(c.value) for c in cells)
            if line.strip():
                print(line)
            n += 1
        if not a.rng and ws.max_row > a.max_rows:
            print(f"... {ws.max_row - a.max_rows} more row(s) — "
                  f"use --range or --max-rows to see them")


if __name__ == "__main__":
    main()
